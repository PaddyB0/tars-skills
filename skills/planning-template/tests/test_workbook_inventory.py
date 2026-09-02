"""Behaviour tests for the thin workbook inventory script.

Seam: the CLI boundary — `python workbook_inventory.py <file.xlsx> [--json]`.
That is how the /planning-template skill invokes it, so it is what we test.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName


SCRIPT = Path(__file__).parents[1] / "workbook_inventory.py"


def run_inventory(
    path: Path, *args: str, env: dict | None = None
) -> subprocess.CompletedProcess:
    child_env = None
    if env is not None:
        child_env = {**os.environ, **env}
    return subprocess.run(
        [sys.executable, str(SCRIPT), str(path), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
        env=child_env,
    )


def run_json(path: Path) -> dict:
    result = run_inventory(path, "--json")
    if result.returncode != 0:
        raise AssertionError(f"exit {result.returncode}: {result.stderr}")
    return json.loads(result.stdout)


class TabListTests(unittest.TestCase):
    def test_lists_every_worksheet_in_workbook_order(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            wb.active.title = "Opex Input"
            wb.create_sheet("Capex Input")
            wb.create_sheet("README")
            wb.save(path)

            result = run_inventory(path)

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Opex Input", result.stdout)
        self.assertIn("Capex Input", result.stdout)
        self.assertIn("README", result.stdout)

    def test_json_tabs_preserve_workbook_order(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            wb.active.title = "First"
            wb.create_sheet("Second")
            wb.create_sheet("Third")
            wb.save(path)

            payload = run_json(path)

        self.assertEqual(
            [tab["name"] for tab in payload["tabs"]],
            ["First", "Second", "Third"],
        )


class HardcodeFormulaRatioTests(unittest.TestCase):
    def test_counts_formula_and_hardcoded_value_cells_per_tab(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Model"
            ws["A1"] = 100
            ws["A2"] = 200
            ws["A3"] = 300
            ws["B1"] = "=A1*2"
            wb.save(path)

            payload = run_json(path)

        tab = payload["tabs"][0]
        self.assertEqual(tab["formula_cells"], 1)
        self.assertEqual(tab["value_cells"], 3)
        self.assertAlmostEqual(tab["hardcode_ratio"], 0.75)

    def test_text_labels_are_not_counted_as_hardcoded_numbers(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Labels"
            ws["A1"] = "Account"
            ws["A2"] = "Rent"
            ws["B2"] = 500
            wb.save(path)

            payload = run_json(path)

        tab = payload["tabs"][0]
        self.assertEqual(tab["value_cells"], 1)
        self.assertEqual(tab["formula_cells"], 0)

    def test_empty_tab_reports_no_ratio_rather_than_dividing_by_zero(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            wb.active.title = "Blank"
            wb.save(path)

            payload = run_json(path)

        tab = payload["tabs"][0]
        self.assertEqual(tab["value_cells"], 0)
        self.assertEqual(tab["formula_cells"], 0)
        self.assertIsNone(tab["hardcode_ratio"])


class NamedRangeTests(unittest.TestCase):
    def test_reports_workbook_scoped_defined_names_with_their_targets(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Data"
            wb.defined_names.add(DefinedName("AccountList", attr_text="Data!$A$1:$A$50"))
            wb.save(path)

            payload = run_json(path)

        names = {n["name"]: n for n in payload["named_ranges"]}
        self.assertIn("AccountList", names)
        self.assertIn("Data!$A$1:$A$50", names["AccountList"]["refers_to"])

    def test_workbook_without_defined_names_reports_an_empty_list(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            wb.active.title = "Solo"
            wb.save(path)

            payload = run_json(path)

        self.assertEqual(payload["named_ranges"], [])


class ExternalLinkTests(unittest.TestCase):
    def test_flags_formula_cells_that_reference_another_workbook(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Opex"
            ws["A1"] = "=[1]Prior!$B$4"
            ws["A2"] = r"='C:\budgets\[FY26.xlsx]Sheet1'!$C$9"
            ws["A3"] = "=A1+1"
            wb.save(path)

            payload = run_json(path)

        links = payload["external_links"]
        cells = {link["cell"] for link in links}
        self.assertEqual(cells, {"Opex!A1", "Opex!A2"})

    def test_workbook_with_no_external_references_reports_an_empty_list(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Clean"
            ws["A1"] = "=SUM(B1:B9)"
            wb.save(path)

            payload = run_json(path)

        self.assertEqual(payload["external_links"], [])


class PerTabDimensionTests(unittest.TestCase):
    def test_reports_header_row_labels_as_that_tabs_dimensions(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Opex Input"
            ws.append(["Department", "Account", "Jan-27", "Feb-27"])
            ws.append(["FACILITIES", "6100 Rent", 1000, 1000])
            wb.save(path)

            payload = run_json(path)

        tab = payload["tabs"][0]
        self.assertEqual(
            tab["dimensions"], ["Department", "Account", "Jan-27", "Feb-27"]
        )

    def test_skips_leading_blank_rows_to_find_the_header(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Padded"
            ws.append([])
            ws.append([])
            ws.append(["Entity", "Clinic"])
            ws.append(["EBH", "SCHERTZ"])
            wb.save(path)

            payload = run_json(path)

        self.assertEqual(payload["tabs"][0]["dimensions"], ["Entity", "Clinic"])

    def test_a_tab_with_no_text_header_reports_no_dimensions(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Numbers"
            ws.append([1, 2, 3])
            wb.save(path)

            payload = run_json(path)

        self.assertEqual(payload["tabs"][0]["dimensions"], [])


class DimensionCandidateTests(unittest.TestCase):
    def test_a_column_of_repeating_labels_is_a_dimension_candidate(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Opex"
            ws.append(["Department", "Account", "Amount"])
            for dept in ["FACILITIES", "MARKETING", "FACILITIES", "IT", "MARKETING"]:
                ws.append([dept, "6100 Rent", 100])
            wb.save(path)

            payload = run_json(path)

        candidates = {c["column"]: c for c in payload["dimension_candidates"]}
        self.assertIn("Department", candidates)
        self.assertEqual(candidates["Department"]["tab"], "Opex")
        self.assertEqual(candidates["Department"]["distinct_values"], 3)
        self.assertEqual(
            sorted(candidates["Department"]["sample_values"]),
            ["FACILITIES", "IT", "MARKETING"],
        )

    def test_a_column_where_every_label_is_unique_is_not_a_candidate(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Opex"
            ws.append(["Note", "Department"])
            for i, dept in enumerate(["A", "B", "A", "B", "A"]):
                ws.append([f"unique note {i}", dept])
            wb.save(path)

            payload = run_json(path)

        columns = {c["column"] for c in payload["dimension_candidates"]}
        self.assertIn("Department", columns)
        self.assertNotIn("Note", columns)

    def test_a_numeric_column_is_not_a_dimension_candidate(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Opex"
            ws.append(["Amount", "Region"])
            for region in ["North", "South", "North", "South", "North"]:
                ws.append([100, region])
            wb.save(path)

            payload = run_json(path)

        columns = {c["column"] for c in payload["dimension_candidates"]}
        self.assertEqual(columns, {"Region"})


class HeuristicDisclosureTests(unittest.TestCase):
    def test_text_output_states_that_candidate_detection_is_heuristic(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Opex"
            ws.append(["Department"])
            for dept in ["A", "B", "A"]:
                ws.append([dept])
            wb.save(path)

            result = run_inventory(path)

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("heuristic", result.stdout.lower())

    def test_json_output_carries_the_same_heuristic_disclosure(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            wb.active.title = "Solo"
            wb.save(path)

            payload = run_json(path)

        self.assertIn("heuristic", payload["dimension_candidates_note"].lower())


class TitleRowTests(unittest.TestCase):
    """Real planning workbooks put a title in row 1 and the header further down."""

    def test_a_single_cell_title_row_does_not_masquerade_as_the_header(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Opex Input"
            ws.append(["2027 Departmental Operating Expense Budget"])
            ws.append([])
            ws.append(["Department", "Account", "Driver method"])
            ws.append(["FACILITIES", "6100 Rent", "Fixed"])
            ws.append(["MARKETING", "6200 Ads", "Fixed"])
            wb.save(path)

            payload = run_json(path)

        self.assertEqual(
            payload["tabs"][0]["dimensions"],
            ["Department", "Account", "Driver method"],
        )

    def test_prose_notes_below_the_data_do_not_become_dimension_candidates(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Ref - Departments"
            ws.append(["Department list"])
            ws.append([])
            ws.append(["Department", "Entity"])
            for dept in ["AUTH", "BILLING", "AUTH", "FINANCE"]:
                ws.append([dept, "EBH"])
            ws.append([])
            ws.append(["Note: DIAGNOSTIC DEPT posts under EBH but is not live."])
            wb.save(path)

            payload = run_json(path)

        columns = {c["column"] for c in payload["dimension_candidates"]}
        self.assertIn("Department", columns)
        self.assertNotIn("Department list", columns)


class ContiguousDataBlockTests(unittest.TestCase):
    def test_column_scan_stops_at_the_first_blank_row_after_the_data(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Ref - Departments"
            ws.append(["Department", "Entity"])
            for dept in ["AUTH", "BILLING", "AUTH", "FINANCE"]:
                ws.append([dept, "EBH"])
            ws.append([])
            ws.append(["Note: DIAGNOSTIC DEPT posts under EBH but is not live."])
            wb.save(path)

            payload = run_json(path)

        dept = next(
            c for c in payload["dimension_candidates"] if c["column"] == "Department"
        )
        self.assertEqual(dept["distinct_values"], 3)
        self.assertEqual(dept["row_count"], 4)
        self.assertEqual(
            sorted(dept["sample_values"]), ["AUTH", "BILLING", "FINANCE"]
        )

    def test_a_blank_gap_inside_the_header_row_does_not_end_the_block(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Gapped"
            ws.append(["Region", "Spacer", "Owner"])
            ws.append(["North", None, "Ana"])
            ws.append(["South", None, "Bo"])
            ws.append(["North", None, "Ana"])
            wb.save(path)

            payload = run_json(path)

        regions = next(
            c for c in payload["dimension_candidates"] if c["column"] == "Region"
        )
        self.assertEqual(regions["row_count"], 3)


class InputErrorTests(unittest.TestCase):
    def test_missing_file_reports_a_readable_error_not_a_traceback(self):
        with tempfile.TemporaryDirectory() as td:
            missing = Path(td) / "absent.xlsx"
            result = run_inventory(missing)

        self.assertEqual(result.returncode, 2)
        self.assertNotIn("Traceback", result.stderr)
        self.assertIn("not found", result.stderr.lower())

    def test_legacy_xls_is_refused_with_the_known_gotcha(self):
        with tempfile.TemporaryDirectory() as td:
            legacy = Path(td) / "budget.xls"
            legacy.write_bytes(b"not really an xls")
            result = run_inventory(legacy)

        self.assertEqual(result.returncode, 2)
        self.assertNotIn("Traceback", result.stderr)
        self.assertIn(".xlsx", result.stderr)

    def test_unreadable_workbook_reports_a_readable_error(self):
        with tempfile.TemporaryDirectory() as td:
            broken = Path(td) / "broken.xlsx"
            broken.write_bytes(b"this is not a zip container")
            result = run_inventory(broken)

        self.assertEqual(result.returncode, 2)
        self.assertNotIn("Traceback", result.stderr)
        self.assertIn("could not read", result.stderr.lower())


class OutputEncodingTests(unittest.TestCase):
    """A narrow console codepage must not decide whether the report renders."""

    def test_report_renders_when_the_console_codepage_is_cp1252(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Data"
            wb.defined_names.add(
                DefinedName("AccountList", attr_text="Data!$A$1:$A$50")
            )
            wb.save(path)

            result = run_inventory(path, env={"PYTHONIOENCODING": "cp1252"})

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertNotIn("UnicodeEncodeError", result.stderr)
        self.assertIn("AccountList", result.stdout)


class ExternalLinksInDefinedNamesTests(unittest.TestCase):
    def test_a_defined_name_pointing_at_another_workbook_is_an_external_link(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            wb.active.title = "Data"
            wb.defined_names.add(
                DefinedName("LegacyGraph", attr_text="'[1]Month 8'!$A$1")
            )
            wb.defined_names.add(
                DefinedName("LocalList", attr_text="Data!$A$1:$A$9")
            )
            wb.save(path)

            payload = run_json(path)

        sources = {link["source"] for link in payload["external_links"]}
        self.assertIn("defined name: LegacyGraph", sources)
        self.assertNotIn("defined name: LocalList", sources)

    def test_cell_sourced_external_links_are_labelled_by_their_cell(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Opex"
            ws["A1"] = "=[1]Prior!$B$4"
            wb.save(path)

            payload = run_json(path)

        self.assertEqual(payload["external_links"][0]["source"], "Opex!A1")


if __name__ == "__main__":
    unittest.main()
