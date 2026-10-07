"""Tests for excel_lint.py — one clean workbook, one rule-by-rule dirty set."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Border, Font, GradientFill, PatternFill, Side
from openpyxl.worksheet.views import Selection

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from excel_lint import lint_workbook  # noqa: E402

INK = "FF1A2233"


def _clean_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    sans = Font(name="Aptos Display", size=10, color=INK)
    title = ws.cell(row=1, column=2, value="Revenue summary")
    title.font = Font(name="Aptos Display", size=14, bold=True, color=INK)
    # mono is for instrument LABELS (letters), not for numeric columns
    ws.cell(row=3, column=2, value="DETAIL").font = Font(
        name="Aptos Mono", size=9, bold=True, color=INK
    )
    for row, value in ((4, 1250.0), (5, -430.0)):
        label = ws.cell(row=row, column=2, value=f"Line {row}")
        label.font = sans
        num = ws.cell(row=row, column=3, value=value)
        num.font = sans          # tabular figures; see DESIGN.md §1
        num.number_format = "#,##0;(#,##0)"
    return wb


def _findings(wb, tmp_path: Path) -> list[dict]:
    path = tmp_path / "book.xlsx"
    wb.save(path)
    return lint_workbook(path)


def _rules(wb, tmp_path: Path) -> set[str]:
    return {f["rule"] for f in _findings(wb, tmp_path)}


def _xl9_details(wb, tmp_path: Path) -> list[str]:
    return [f["detail"] for f in _findings(wb, tmp_path) if f["rule"] == "XL9"]


def _excel_two_axis_freeze(ws, cursor: str | None) -> None:
    """Freeze at B8 the way Excel 365 saves it (verified against Excel output).

    One selection per pane, topRight first. The active pane (bottomRight)
    holds the real cursor; Excel omits its activeCell when the cursor is A1.
    """
    ws.freeze_panes = "B8"
    ws.sheet_view.pane.activePane = "bottomRight"
    ws.sheet_view.selection = [
        Selection(pane="topRight", activeCell="B1", sqref="B1"),
        Selection(pane="bottomLeft", activeCell="A8", sqref="A8"),
        Selection(pane="bottomRight", activeCell=cursor, sqref=cursor),
    ]


def test_clean_workbook_has_no_findings(tmp_path: Path) -> None:
    assert _rules(_clean_workbook(), tmp_path) == set()


def test_default_workbook_fires_gridlines_tabname_and_general(tmp_path: Path) -> None:
    wb = Workbook()
    wb.active.cell(row=1, column=1, value=42.5)
    rules = _rules(wb, tmp_path)
    assert {"XL1", "XL3", "XL7"} <= rules


def test_merged_cells_fire_xl2(tmp_path: Path) -> None:
    wb = _clean_workbook()
    wb["Summary"].merge_cells("B7:D7")
    assert "XL2" in _rules(wb, tmp_path)


def test_font_sprawl_fires_xl4(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    ws.cell(row=8, column=2, value="stray").font = Font(name="Comic Sans MS", size=10)
    assert "XL4" in _rules(wb, tmp_path)


def test_four_sizes_are_clean_but_a_fifth_fires_xl4(tmp_path: Path) -> None:
    """A style has four sizes; chart titles reuse the body size."""
    wb = _clean_workbook()
    ws = wb["Summary"]
    ws.cell(row=8, column=2, value="kpi").font = Font(name="Aptos Display", size=16, color=INK)
    assert "XL4" not in _rules(wb, tmp_path)
    ws.cell(row=9, column=2, value="chart title").font = Font(
        name="Aptos Display", size=11, bold=True, color=INK
    )
    assert "XL4" in _rules(wb, tmp_path)


def test_conflicting_shared_edge_fires_xl17(tmp_path: Path) -> None:
    """Two stacked cells each declare their shared edge, in different colors."""
    wb = _clean_workbook()
    ws = wb["Summary"]
    ws["H8"].border = Border(bottom=Side(style="thin", color="FF8F5F17"))
    ws["H9"].border = Border(top=Side(style="thin", color="FF1C593E"))
    assert "XL17" in _rules(wb, tmp_path)


def test_shared_edge_declared_once_or_identically_is_clean(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    ws["H8"].border = Border(bottom=Side(style="thin", color="FFD5D7DF"))
    ws["H9"].border = Border(top=Side(style="thin", color="FFD5D7DF"), right=Side(style="thin"))
    ws["C8"].border = Border(right=Side(style="medium", color="FF202B8A"))
    assert "XL17" not in _rules(wb, tmp_path)


def _period_sheet(columns: int):
    """A landscape Letter sheet, fit to one page wide, 0.5in side margins."""
    from openpyxl.worksheet.properties import PageSetupProperties

    wb = _clean_workbook()
    ws = wb["Summary"]
    ws.column_dimensions["A"].width = 2.29
    ws.column_dimensions["B"].width = 28.71
    last = 2 + columns
    for col in range(3, last + 1):
        letter = ws.cell(row=4, column=col).column_letter
        ws.column_dimensions[letter].width = 14
        cell = ws.cell(row=4, column=col, value=1000.0 + col)
        cell.font = Font(name="Aptos Display", size=10, color=INK)
        cell.number_format = "#,##0;(#,##0)"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.print_area = f"A1:{ws.cell(row=5, column=last).coordinate}"
    return wb


def test_wide_fit_to_width_sheet_fires_xl18(tmp_path: Path) -> None:
    """Sixteen 14-wide columns shrink 10pt numbers to about 5pt on Letter landscape."""
    assert "XL18" in _rules(_period_sheet(16), tmp_path)


def test_sheet_that_fits_at_full_scale_is_clean_for_xl18(tmp_path: Path) -> None:
    assert "XL18" not in _rules(_period_sheet(4), tmp_path)


def test_fit_to_one_page_tall_counts_toward_xl18(tmp_path: Path) -> None:
    """A narrow sheet that passes on width still shrinks when forced onto one page tall."""
    wb = _period_sheet(4)
    ws = wb["Summary"]
    for row in range(6, 206):
        cell = ws.cell(row=row, column=3, value=float(row))
        cell.font = Font(name="Aptos Display", size=10, color=INK)
        cell.number_format = "#,##0;(#,##0)"
    ws.page_setup.fitToHeight = 1
    ws.print_area = "A1:G205"
    assert "XL18" in _rules(wb, tmp_path)


def test_missing_fit_attributes_mean_one_page(tmp_path: Path) -> None:
    """Excel reads an absent fitToWidth as 1 page, so a wide sheet still shrinks."""
    wb = _period_sheet(16)
    ws = wb["Summary"]
    ws.page_setup.fitToWidth = None
    ws.page_setup.fitToHeight = None
    assert "XL18" in _rules(wb, tmp_path)


def test_small_text_on_a_shrunk_sheet_fires_xl18(tmp_path: Path) -> None:
    """The floor covers every printed line, not only numbers: an 8pt note at 80% prints at 6.4pt."""
    wb = _period_sheet(9)
    assert "XL18" not in _rules(wb, tmp_path)
    ws = wb["Summary"]
    ws.cell(row=5, column=2, value="Variance is favorability, not sign.").font = Font(
        name="Aptos Display", size=8, italic=True, color=INK
    )
    assert "XL18" in _rules(wb, tmp_path)


def test_fixed_print_scale_below_floor_fires_xl18(tmp_path: Path) -> None:
    wb = _clean_workbook()
    wb["Summary"].page_setup.scale = 60  # 10pt numbers print at 6pt
    assert "XL18" in _rules(wb, tmp_path)


def test_pure_black_fires_xl5(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    ws.cell(row=8, column=2, value="black").font = Font(
        name="Calibri", size=10, color="FF000000"
    )
    assert "XL5" in _rules(wb, tmp_path)


def test_fill_sprawl_fires_xl6(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    hexes = ["FFFF0000", "FF00FF00", "FF0000FF", "FFFFFF00", "FFFF00FF", "FF00FFFF", "FF884400"]
    for i, rgb in enumerate(hexes):
        cell = ws.cell(row=10 + i, column=2, value=f"fill {i}")
        cell.font = Font(name="Calibri", size=10, color=INK)
        cell.fill = PatternFill(patternType="solid", fgColor=rgb)
    assert "XL6" in _rules(wb, tmp_path)


def test_long_table_without_freeze_fires_xl8(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    font = Font(name="Calibri", size=10, color=INK)
    for row in range(4, 40):
        cell = ws.cell(row=row, column=2, value=f"row {row}")
        cell.font = font
    assert "XL8" in _rules(wb, tmp_path)
    ws.freeze_panes = "B5"
    assert "XL8" not in _rules(wb, tmp_path)


def test_unparked_sheet_fires_xl9(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    ws.sheet_view.zoomScale = 85
    assert "XL9" in _rules(wb, tmp_path)


def test_two_axis_freeze_parked_at_a1_is_clean_for_xl9(tmp_path: Path) -> None:
    """The topRight pane starts at column B, so its B1 is not the cursor."""
    wb = _clean_workbook()
    _excel_two_axis_freeze(wb["Summary"], cursor=None)
    assert _xl9_details(wb, tmp_path) == []


def test_two_axis_freeze_cursor_elsewhere_fires_xl9_at_the_cursor(tmp_path: Path) -> None:
    wb = _clean_workbook()
    _excel_two_axis_freeze(wb["Summary"], cursor="D20")
    assert _xl9_details(wb, tmp_path) == ["cursor parked at D20; park at A1"]


def test_row_only_freeze_parked_at_a1_is_clean_for_xl9(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    ws.freeze_panes = "A8"
    ws.sheet_view.selection = [Selection(pane="bottomLeft", activeCell=None, sqref=None)]
    assert _xl9_details(wb, tmp_path) == []


def test_unfrozen_sheet_parked_elsewhere_fires_xl9(tmp_path: Path) -> None:
    wb = _clean_workbook()
    wb["Summary"].sheet_view.selection = [Selection(activeCell="C5", sqref="C5")]
    assert _xl9_details(wb, tmp_path) == ["cursor parked at C5; park at A1"]


def test_tab_rainbow_fires_xl10(tmp_path: Path) -> None:
    wb = _clean_workbook()
    for i, rgb in enumerate(["FFFF0000", "FF00FF00", "FF0000FF", "FFFFFF00"]):
        ws = wb.create_sheet(f"Tab {i}")
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = rgb
    assert "XL10" in _rules(wb, tmp_path)


def test_border_grid_fires_xl12(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    side = Side(style="thin")
    boxed = Border(top=side, bottom=side, left=side, right=side)
    font = Font(name="Calibri", size=10, color=INK)
    for row in range(4, 40):
        cell = ws.cell(row=row, column=2, value=f"cell {row}")
        cell.font = font
        cell.border = boxed
    assert "XL12" in _rules(wb, tmp_path)


def _boxed_column(wb, sides: dict) -> None:
    ws = wb["Summary"]
    font = Font(name="Aptos Display", size=10, color="FF1D59A3")
    fill = PatternFill("solid", fgColor="FFD6E4F5")
    for row in range(4, 40):
        cell = ws.cell(row=row, column=3, value=row * 10)
        cell.font = font
        cell.fill = fill
        cell.number_format = "#,##0;(#,##0)"
        cell.border = Border(**sides)


def test_dashed_input_boxes_are_exempt_from_xl12(tmp_path: Path) -> None:
    dash = Side(style="dashed", color="FF63677E")
    wb = _clean_workbook()
    _boxed_column(wb, dict(top=dash, bottom=dash, left=dash, right=dash))
    assert "XL12" not in _rules(wb, tmp_path)


def test_input_box_that_meets_a_structural_rule_stays_exempt(tmp_path: Path) -> None:
    dash = Side(style="dashed", color="FF63677E")
    wb = _clean_workbook()
    _boxed_column(wb, dict(top=dash, bottom=Side(style="thin", color="FF292B36"), left=dash, right=dash))
    assert "XL12" not in _rules(wb, tmp_path)


def test_two_dashed_sides_still_count_as_a_border_grid(tmp_path: Path) -> None:
    dash = Side(style="dashed", color="FF63677E")
    thin = Side(style="thin", color="FFD5D7DF")
    wb = _clean_workbook()
    _boxed_column(wb, dict(top=thin, bottom=thin, left=dash, right=dash))
    assert "XL12" in _rules(wb, tmp_path)


def test_hidden_sheets_skip_view_rules(tmp_path: Path) -> None:
    wb = _clean_workbook()
    staging = wb.create_sheet("Staging")
    staging.cell(row=1, column=1, value=99.9)
    staging.sheet_state = "hidden"
    rules = _rules(wb, tmp_path)
    assert "XL1" not in rules
    assert "XL7" not in rules


def test_sans_mono_pair_is_clean_but_third_face_fires_xl4_and_xl15(tmp_path: Path) -> None:
    wb = _clean_workbook()
    assert _rules(wb, tmp_path) == set()
    ws = wb["Summary"]
    ws.cell(row=8, column=2, value="stray").font = Font(name="Verdana", size=10, color=INK)
    rules = _rules(wb, tmp_path)
    assert "XL4" in rules
    assert "XL15" in rules


def test_unapproved_pair_fires_xl15_without_xl4(tmp_path: Path) -> None:
    """Two faces is within budget, but Comic Sans is not an approved pair."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=2, value="Title").font = Font(
        name="Comic Sans MS", size=14, color=INK
    )
    n = ws.cell(row=4, column=3, value=1.0)
    n.font = Font(name="Papyrus", size=10, color=INK)
    n.number_format = "#,##0"
    rules = _rules(wb, tmp_path)
    assert "XL15" in rules
    assert "XL4" not in rules


def test_gradient_fill_fires_xl13(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    cell = ws.cell(row=8, column=2, value="shiny")
    cell.font = Font(name="Calibri", size=10, color=INK)
    cell.fill = GradientFill(stop=("FFFFFFFF", "FF202B8A"))
    assert "XL13" in _rules(wb, tmp_path)


def test_data_bar_fires_xl14(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    ws.conditional_formatting.add(
        "C4:C5", DataBarRule(start_type="min", end_type="max", color="FF202B8A")
    )
    assert "XL14" in _rules(wb, tmp_path)


def test_color_scale_fires_xl14(tmp_path: Path) -> None:
    wb = _clean_workbook()
    ws = wb["Summary"]
    ws.conditional_formatting.add(
        "C4:C5",
        ColorScaleRule(
            start_type="min", start_color="FFFFFFFF", end_type="max", end_color="FF202B8A"
        ),
    )
    assert "XL14" in _rules(wb, tmp_path)


def test_proportional_figure_numerals_fire_xl16(tmp_path: Path) -> None:
    """Inter's digits are not tabular (measured 16.0-26.0 units), so a numeric
    column set in it can never align. Text in Inter is not the defect."""
    wb = _clean_workbook()
    ws = wb["Summary"]
    for row in (4, 5):
        ws.cell(row=row, column=3).font = Font(name="Inter", size=10, color=INK)
    assert "XL16" in _rules(wb, tmp_path)


def test_tabular_figure_numerals_do_not_fire_xl16(tmp_path: Path) -> None:
    assert "XL16" not in _rules(_clean_workbook(), tmp_path)


def test_mono_numerals_do_not_fire_xl16(tmp_path: Path) -> None:
    """Consolas is tabular too - it is permitted, just not preferred for
    dense columns (31% wider, slashed zero). XL16 is about alignment only."""
    wb = _clean_workbook()
    ws = wb["Summary"]
    for row in (4, 5):
        ws.cell(row=row, column=3).font = Font(name="Consolas", size=10, color=INK)
    assert "XL16" not in _rules(wb, tmp_path)


def test_inter_numerals_also_fire_xl15_as_unapproved(tmp_path: Path) -> None:
    wb = _clean_workbook()
    wb["Summary"].cell(row=4, column=3).font = Font(name="Inter", size=10, color=INK)
    assert "XL15" in _rules(wb, tmp_path)


def test_calibri_consolas_fallback_pair_is_approved(tmp_path: Path) -> None:
    """The pre-2024-Office fallback pair stays valid on its own."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.cell(row=3, column=2, value="DETAIL").font = Font(
        name="Consolas", size=9, bold=True, color=INK
    )
    n = ws.cell(row=4, column=3, value=1250.0)
    n.font = Font(name="Calibri", size=10, color=INK)
    n.number_format = "#,##0"
    assert _rules(wb, tmp_path) == set()


def test_mixing_two_approved_pairs_fires_xl15(tmp_path: Path) -> None:
    """Aptos + Consolas is two faces from different families - not a pair."""
    wb = _clean_workbook()
    wb["Summary"].cell(row=8, column=2, value="LABEL").font = Font(
        name="Consolas", size=9, bold=True, color=INK
    )
    rules = _rules(wb, tmp_path)
    assert "XL15" in rules


def test_aptos_regular_pairs_with_aptos_mono(tmp_path: Path) -> None:
    """Aptos regular stays an approved alternate to Aptos Display."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.cell(row=3, column=2, value="DETAIL").font = Font(
        name="Aptos Mono", size=9, bold=True, color=INK
    )
    n = ws.cell(row=4, column=3, value=1250.0)
    n.font = Font(name="Aptos", size=10, color=INK)
    n.number_format = "#,##0"
    assert _rules(wb, tmp_path) == set()


def test_mixing_display_with_aptos_regular_fires_xl15(tmp_path: Path) -> None:
    """Both are Aptos cuts but they are different Excel families - one pair."""
    wb = _clean_workbook()
    wb["Summary"].cell(row=8, column=2, value="stray").font = Font(
        name="Aptos", size=10, color=INK
    )
    assert "XL15" in _rules(wb, tmp_path)


def test_aptos_narrow_pairs_with_aptos_mono(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.cell(row=3, column=2, value="DETAIL").font = Font(
        name="Aptos Mono", size=9, bold=True, color=INK
    )
    n = ws.cell(row=4, column=3, value=1250.0)
    n.font = Font(name="Aptos Narrow", size=10, color=INK)
    n.number_format = "#,##0"
    assert _rules(wb, tmp_path) == set()


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
